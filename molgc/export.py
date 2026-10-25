"""Exportación de resultados a Excel (.xlsx) con openpyxl."""
from __future__ import annotations

import io
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F3A8A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="E8EDF7")
THIN = Side(style="thin", color="C9D1DE")


def _sheet(wb: Workbook, title: str, header: List[str], rows: List[List], widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    import re as _re
    ws.title = _re.sub(r"[\[\]:*?/\\]", "_", title)[:31]
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(i) or min(max(10, len(str(h)) * 0.9), 42)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 36
    return ws


def _fmt(ws, start_col: int, end_col: int, fmt: str):
    for row in ws.iter_rows(min_row=2, min_col=start_col, max_col=end_col):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = fmt


def comparison_workbook(res: Dict, title: str = "MolGC") -> bytes:
    wb = Workbook()
    labels = res["labels"]
    ref = res["reference"]["label"]
    # Hoja 1: resumen
    head = ["Geometría evaluada", "Software", "Funcional detectado", "Enlaces", "Enlaces emparejados",
            "MAE (Å)", "MAPE (%)", "Error abs. máx. (Å)", "RMSE enlaces (Å)", "RMSD ICP (Å)",
            "RMSD con correspondencia (Å)", "Enlaces verificados"]
    rows = []
    for r in res["results"]:
        lc = r["label_check"]
        ver = f"{lc['atoms_ok']}/{lc['atoms_checked']}" if lc["atoms_checked"] else "n/d"
        rows.append([r["label"], r["software"], r["functional"] or "", r["n_bonds"], r["n_matched"], r["MAE"],
                     r["MAPE"], r["max_abs_error"], r["RMSE_bonds"], r["RMSD_ICP"], r["RMSD_matched"], ver])
    ws = _sheet(wb, "Resumen", head, rows, {1: 30}, first=True)
    _fmt(ws, 6, 11, "0.000000")
    ws.append([])
    ws.append([f"Referencia: {ref}  |  {res['reference']['formula']}  |  {res['reference']['n_bonds']} enlaces"])
    o = res["options"]
    ws.append([f"Criterio de enlace: {o['bond_mode']}  |  Emparejamiento: {o['matching']}"])
    ws.append(["MAE = (1/|E|) Σ |L_B(g) − L_B(ref)|   (ec. 33 del artículo).  MAPE = MAE relativo a la longitud de referencia."])
    ws.append(["Cite: Camarillo-Cisneros J. et al. MolGC: molecular geometry comparator algorithm for bond length "
               "mean absolute error computation on molecules. Mol. Divers. (2024). doi:10.1007/s11030-024-10945-2"])

    # Hoja 2: formato de la Tabla 2 del artículo
    head = ["Etiqueta B", "Átomo 1", "Átomo 2", "Índices (ref)", f"Ref: {ref} (Å)"]
    head += [f"{l} (Å)" for l in labels] + [f"|{l} − ref| (Å)" for l in labels]
    rows = []
    for b in res["bonds"]:
        row = [b["B"], b["atom1"], b["atom2"], f"{b['i']}-{b['j']}", b["ref"]]
        row += [b.get(f"len::{l}") for l in labels]
        row += [None if b.get(f"err::{l}") is None else abs(b[f"err::{l}"]) for l in labels]
        rows.append(row)
    mae_row = ["MAE", "", "", "", ""] + [""] * len(labels) + [r["MAE"] for r in res["results"]]
    ws = _sheet(wb, "Comparación y MAE", head, rows + [mae_row], {4: 12})
    _fmt(ws, 5, len(head), "0.00000")
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = SUB_FILL

    # Hoja 3: MAE por tipo de enlace
    pairs = sorted({p for r in res["results"] for p in r["per_pair"]})
    head = ["Tipo de enlace"] + [f"MAE {l} (Å)" for l in labels] + ["n"]
    rows = []
    for p in pairs:
        n = next((r["per_pair"][p]["n"] for r in res["results"] if p in r["per_pair"]), "")
        rows.append([p] + [r["per_pair"].get(p, {}).get("MAE") for r in res["results"]] + [n])
    ws = _sheet(wb, "MAE por tipo de enlace", head, rows)
    _fmt(ws, 2, len(head) - 1, "0.000000")

    # Hoja 4: descriptores
    rows = [[b["B"], b["pair"], b["descriptor"]] for b in res["bonds"]]
    _sheet(wb, "Descriptores", ["Etiqueta B", "Enlace", "Descriptor de cadena S_uv (referencia)"], rows, {3: 90})

    if res.get("warnings"):
        _sheet(wb, "Avisos", ["Aviso"], [[w] for w in res["warnings"]], {1: 120})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def study_workbook(st: Dict) -> bytes:
    wb = Workbook()
    F = st["functionals"]
    head = ["Familia", "Molécula", "Enlaces"] + [f"MAE {f} (Å)" for f in F]
    rows = [[r["family"], r["molecule"], r["n_bonds"]] + [r.get(f) for f in F] for r in st["per_molecule"]]
    ws = _sheet(wb, "MAE por molécula", head, rows, {1: 20, 2: 22}, first=True)
    _fmt(ws, 4, len(head), "0.000000")
    ws.append([])
    ws.append([f"Referencia: {st['reference']}"])

    head = ["Familia", "Molécula"] + [f"RMSD ICP {f} (Å)" for f in F]
    rows = [[r["family"], r["molecule"]] + [r.get(f"RMSD::{f}") for f in F] for r in st["per_molecule"]]
    ws = _sheet(wb, "RMSD por molécula", head, rows, {1: 20, 2: 22})
    _fmt(ws, 3, len(head), "0.00000")

    head = ["Funcional"] + st["families"] + ["Promedio"]
    rows = [[f] + [st["family_mae"][fam].get(f) for fam in st["families"]] + [st["overall_mae"].get(f)] for f in F]
    ws = _sheet(wb, "MAE por familia", head, rows, {1: 26})
    _fmt(ws, 2, len(head), "0.000000")

    if st["topsis"]:
        head = ["Ranking", "Funcional", "Cercanía relativa C*", "Distancia a la ideal", "Distancia a la anti-ideal"]
        rows = [[t["rank"], t["functional"], t["closeness"], t["d_ideal"], t["d_anti_ideal"]] for t in st["topsis"]]
        ws = _sheet(wb, "UW-TOPSIS", head, rows, {2: 26})
        _fmt(ws, 3, 5, "0.0000")
        ws.append([])
        ws.append(["Alternativas: funcionales; criterios: familias (sin pesos). C* = d⁻/(d⁺+d⁻); 1 = solución ideal."])

    for m in st["molecules"]:
        labels = m["labels"]
        head = ["B", "Átomo 1", "Átomo 2", f"Ref (Å)"] + [f"{l} (Å)" for l in labels]
        rows = [[b["B"], b["atom1"], b["atom2"], b["ref"]] + [b.get(f"len::{l}") for l in labels] for b in m["bonds"]]
        rows.append(["MAE", "", "", ""] + [r["MAE"] for r in m["results"]])
        ws = _sheet(wb, f"{m['molecule']}"[:31], head, rows)
        _fmt(ws, 4, len(head), "0.00000")
    if st.get("warnings"):
        _sheet(wb, "Avisos", ["Aviso"], [[w] for w in st["warnings"]], {1: 120})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
